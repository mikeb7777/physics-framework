# -*- coding: utf-8 -*-
"""Version 6F of A Quest for The Big TOE, from Version 6E (Michael, 10 Oct 2026: "We can do 6F now").

Applies the changes in 6F-PROPOSED-CHANGES.md, which came out of the global research audit
(GLOBAL-RESEARCH-AUDIT-2026-10.md):

  A1  Introduction, Fibonacci: credit the Indian prosodists                      new ref Intro [59]
  A2  Element 1, pi: Egypt, Archimedes, China, Madhava, al-Kashi                  new refs E1 [30]-[33]
  A3  Introduction, binary: Egyptian multiplication by doubling                   new ref Intro [60]
  A4  Element 2, first use of "algorithm": al-Khwarizmi                           new ref E2 [28]
  B1  Element 21: USTC's independent below-threshold result (He et al., PRL 2025)  new ref E21 [12]
  --  Edition line: "Version 6C" (left unchanged through 6D and 6E) -> "Version 6F"

The book's references live on references.html (each Element ends "See: .../references.html"), so
the new references are added there with these numbers, in the same commit.

New runs copy the formatting of the run they follow; new paragraphs are copies of the paragraph
they follow, with the text replaced. Run once:  python3 tools/book_6f_2026_10_10.py
"""
import copy, os
import docx

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "A Quest for The Big TOE Version 6E.docx")
DST = os.path.join(ROOT, "A Quest for The Big TOE Version 6F.docx")

d = docx.Document(SRC)
P = d.paragraphs


def para(start):
    """The one paragraph whose text starts with `start`."""
    hits = [p for p in P if p.text.startswith(start)]
    assert len(hits) == 1, (start, len(hits))
    return hits[0]


def insert_in_run(p, after, text):
    """Insert `text` straight after the substring `after`, inside the run that holds it."""
    runs = [r for r in p.runs if after in r.text]
    assert len(runs) == 1, (after, len(runs))
    r = runs[0]
    assert r.text.count(after) == 1, after
    r.text = r.text.replace(after, after + text)


def add_paragraph_after(p, text):
    """A copy of paragraph `p` (style and run formatting) holding `text`, placed after it."""
    new = copy.deepcopy(p._p)
    p._p.addnext(new)
    q = docx.text.paragraph.Paragraph(new, p._parent)
    for r in q.runs[1:]:
        r._r.getparent().remove(r._r)
    q.runs[0].text = text
    return q


# Edition line
ed = [p for p in P if p.text.strip() == "Version 6C"]
assert len(ed) == 1
ed[0].runs[0].text = "Version 6F"

# A1  Fibonacci
p = para("This is the Fibonacci sequence: 1, 1, 2, 3, 5, 8, 13.")
insert_in_run(p, "This is the Fibonacci sequence: 1, 1, 2, 3, 5, 8, 13.",
              " It carries the name of Leonardo of Pisa, who brought it to Europe in 1202, but Indian"
              " scholars of poetic meter had described it centuries earlier, from Pingala to Virahanka"
              " and Hemachandra [59].")

# A3  binary: a new paragraph after "Every binary in physics traces back..."
p = para("Every binary in physics traces back to this original twoness.")
add_paragraph_after(p,
    "People found that twoness long before physics did. Egyptian scribes multiplied by doubling: to"
    " multiply by 23, they doubled the other number four times and added the rows for 16, 4, 2 and 1,"
    " which is writing 23 in binary. The Rhind papyrus, copied around 1550 BCE from an older text,"
    " records the method, and a halving-and-doubling form of it is still taught in Ethiopia and"
    " Eritrea [60].")

# A2  pi: after the [13] citation run in Element 1
p = para("Mathematical constants reveal the same secret operating at a deeper level. Pi doesn’t")
runs = p.runs
i13 = [k for k, r in enumerate(runs) if r.text == "[13]"]
assert len(i13) == 1
nxt = runs[i13[0] + 1]
assert nxt.text.startswith(". The golden ratio phi"), nxt.text[:40]
nxt.text = (". That relationship was pursued independently across the ancient and medieval world."
            " Egyptian scribes worked with a value close to 3.16 in a text first written about 1850 BCE."
            " Archimedes bounded it with polygons. Liu Hui and Zu Chongzhi in China carried the polygon"
            " method to seven decimal places by the fifth century. Madhava in Kerala found an infinite"
            " series for it around 1400, and in 1424 al-Kashi in Samarkand computed it to sixteen decimal"
            " places, a record that stood for nearly two centuries [30, 31, 32, 33]"
            + nxt.text)

# A4  algorithm
p = para("Reversible computing offers a related path. By designing algorithms that avoid")
insert_in_run(p, "By designing algorithms",
              " (the word comes from al-Khwarizmi, the ninth-century Persian scholar in Baghdad whose name"
              " was Latinized as Algoritmi, and whose book on equations also gave us “algebra” [28])")

# B1  Element 21: after "After 30 years, the threshold had finally been crossed."
p = para("Exponential error suppression means you can keep adding qubits and errors keep decreasing")
assert p.text.rstrip().endswith("After 30 years, the threshold had finally been crossed."), p.text[-80:]
add_paragraph_after(p,
    "A single result is a claim; a second, independent one is the beginning of knowledge. In December"
    " 2025 a team at the University of Science and Technology of China reported the same threshold"
    " crossing on its own processor, Zuchongzhi 3.2, with a distance-7 surface code [12]. Its"
    " suppression factor, about 1.4 for each step up in code size, was smaller than Willow’s 2.14,"
    " but it was above one, built on different hardware, and achieved by a different group, using"
    " microwave control alone to stop qubits leaking out of their two working states. The threshold"
    " is no longer one laboratory’s result.")

d.core_properties.revision = (d.core_properties.revision or 0) + 1
d.save(DST)
print("wrote", os.path.basename(DST))
