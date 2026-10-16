# -*- coding: utf-8 -*-
"""Version 6D, new cover (Michael, 16 Oct 2026: "Please replace the book cover and back on 6D with the new cover image").

The front-cover picture anchored to page 1 is replaced with the new front (direction A from the cover
canvas, 6 x 9 in at 300 DPI), and the new back cover is added as the last page, anchored the same way.
Both are pinned to the page's top-left corner at exactly 6 x 9 in. Images: cover/front.jpg, cover/back.jpg.
Run once:  python3 tools/book_6d_new_cover_2026_10_16.py
"""
import copy, os, re
import docx
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC = os.path.join(ROOT, "A Quest for The Big TOE Version 6D.docx")
FRONT = os.path.join(ROOT, "cover", "front.jpg")
BACK = os.path.join(ROOT, "cover", "back.jpg")
EMU_IN = 914400

d = docx.Document(DOC)
body = d.element.body

# the anchored full-page picture on page 1 (the only anchor with a 6 in wide extent before the title)
anchors = [a for a in body.iter(qn("wp:anchor")) if int(a.find(qn("wp:extent")).get("cx")) == 5486484]
assert len(anchors) == 1, len(anchors)
front_anchor = anchors[0]


def pin_full_page(anchor, rid, name):
    for tag in ("wp:positionH", "wp:positionV"):
        pos = anchor.find(qn(tag))
        pos.set("relativeFrom", "page")
        pos.find(qn("wp:posOffset")).text = "0"
    ext = anchor.find(qn("wp:extent")); ext.set("cx", str(6 * EMU_IN)); ext.set("cy", str(9 * EMU_IN))
    aext = anchor.find(".//" + qn("a:ext")); aext.set("cx", str(6 * EMU_IN)); aext.set("cy", str(9 * EMU_IN))
    anchor.find(".//" + qn("a:blip")).set(qn("r:embed"), rid)
    anchor.find(qn("wp:docPr")).set("name", name)
    anchor.find(qn("wp:docPr")).set("descr", name)


# front: point the picture at the new image
rid_front, _ = d.part.get_or_add_image(FRONT)
pin_full_page(front_anchor, rid_front, "Front cover")

# back: a new last page holding a copy of the front's drawing, pointed at the back image
rid_back, _ = d.part.get_or_add_image(BACK)
drawing_run = copy.deepcopy(front_anchor.getparent().getparent())    # the <w:r> holding <w:drawing>
assert drawing_run.tag == qn("w:r")
back_anchor = drawing_run.find(".//" + qn("wp:anchor"))
pin_full_page(back_anchor, rid_back, "Back cover")
ids = [int(x.get("id")) for x in body.iter(qn("wp:docPr"))]
back_anchor.find(qn("wp:docPr")).set("id", str(max(ids) + 1))
back_anchor.set("relativeHeight", str(int(back_anchor.get("relativeHeight")) + 1))

last = d.paragraphs[-1]
p = d.add_paragraph()                      # python-docx places it before the final sectPr
p.style = last.style
p.add_run().add_break(WD_BREAK.PAGE)
p2 = d.add_paragraph(); p2.style = last.style
p2._p.append(drawing_run)

d.save(DOC)
print("front", rid_front, "back", rid_back)

# the old front image is no longer referenced: drop it
d = docx.Document(DOC)
for rid, rel in list(d.part.rels.items()):
    if rel.reltype.endswith("/image") and rel.target_ref == "media/image1.jpeg" and f'"{rid}"' not in d.element.xml:
        d.part.drop_rel(rid)
d.save(DOC)
